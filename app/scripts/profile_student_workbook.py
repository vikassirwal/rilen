from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.datetime import from_excel


NULL_TOKENS = {"", "?", "NA", "N/A", "#REF!"}


def clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        cleaned = re.sub(r"\s+", " ", value.replace("\xa0", " ")).strip()
        return None if cleaned.upper() in NULL_TOKENS else cleaned
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def normalize_header(value: Any) -> str | None:
    cleaned = clean_cell(value)
    if cleaned is None:
        return None
    return re.sub(r"\s+", " ", str(cleaned).upper()).strip()


def row_has_data(row: tuple[Any, ...]) -> bool:
    return any(clean_cell(value) is not None for value in row)


def is_header_row(row: tuple[Any, ...]) -> bool:
    return normalize_header(row[0]) in {"S.NO", "S. NO", "SR.NO", "SR. NO"}


def section_marker(row: tuple[Any, ...]) -> str | None:
    joined = " ".join(str(clean_cell(value) or "").upper() for value in row)
    compact = re.sub(r"[^A-Z0-9]+", "", joined)
    if compact == "RTE" or ("RTE" in compact and "STUDENT" in compact):
        return "rte"
    if "LEFT" in joined:
        return "left"
    return None


def parse_date(value: Any) -> str | None:
    cleaned = clean_cell(value)
    if cleaned is None:
        return None
    if isinstance(cleaned, datetime):
        return cleaned.date().isoformat()
    if isinstance(cleaned, (int, float)):
        try:
            return from_excel(cleaned).date().isoformat()
        except Exception:
            return None
    if isinstance(cleaned, str):
        for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(cleaned, fmt).date().isoformat()
            except ValueError:
                continue
    return None


def main() -> None:
    workbook_path = Path(sys.argv[1])
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    report: dict[str, Any] = {
        "workbook": str(workbook_path),
        "sheets": [],
    }

    for worksheet in workbook.worksheets:
        section = "enrolled"
        headers: list[str | None] = []
        counts: Counter[str] = Counter()
        section_headers: dict[str, list[str | None]] = {}
        samples: dict[str, list[dict[str, Any]]] = defaultdict(list)
        date_issues: list[dict[str, Any]] = []
        markers: list[dict[str, Any]] = []

        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            marker = section_marker(row)
            if marker:
                section = marker
                markers.append({"row": row_number, "section": marker})
                continue

            if is_header_row(row):
                headers = [normalize_header(value) for value in row]
                section_headers[section] = headers
                continue

            if not headers or not row_has_data(row):
                continue

            row_data = {}
            for index, header in enumerate(headers):
                if not header or header in {"S.NO", "S. NO", "SR.NO", "SR. NO"}:
                    continue
                value = clean_cell(row[index] if index < len(row) else None)
                row_data[header] = value

                if header in {"DOB", "ADMN.DATE", "ADM DATE", "ADMISSION DATE"} and value is not None:
                    if parse_date(value) is None:
                        date_issues.append({"row": row_number, "header": header, "value": value})

            name = row_data.get("NAME")
            surname = row_data.get("SURNAME")
            if clean_cell(name) is None and clean_cell(surname) is None:
                continue

            counts[section] += 1
            if len(samples[section]) < 3:
                samples[section].append({"row": row_number, "data": row_data})

        report["sheets"].append(
            {
                "sheet": worksheet.title,
                "dimensions": {"rows": worksheet.max_row, "columns": worksheet.max_column},
                "markers": markers,
                "counts": dict(counts),
                "headers_by_section": section_headers,
                "samples": samples,
                "date_issues": date_issues[:25],
                "date_issue_count": len(date_issues),
            }
        )

    print(json.dumps(report, indent=2, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
